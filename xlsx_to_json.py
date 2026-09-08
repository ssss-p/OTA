import json
import os

import openpyxl


def xlsx_to_json(xlsx_path, output_dir):
    """把 Excel 的每个 sheet 转换成一个 JSON 文件。"""
    os.makedirs(output_dir, exist_ok=True)
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)

    result = {}
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        headers = [cell.value for cell in ws[1]]
        rows = []

        for row in ws.iter_rows(min_row=2, values_only=True):
            row_data = {}
            for header, value in zip(headers, row):
                if header is None:
                    continue
                row_data[header] = value
            rows.append(row_data)

        result[sheet_name] = rows

        json_name = f"{sheet_name}.json"
        json_path = os.path.join(output_dir, json_name)
        with open(json_path, "w", encoding="utf-8") as f:
            json.dump(rows, f, ensure_ascii=False, indent=2)
        print(f"已生成: {json_path}")

    wb.close()
    return result


if __name__ == "__main__":
    xlsx_file = "OTA.xlsx"
    output_dir = "ota_json"
    xlsx_to_json(xlsx_file, output_dir)
